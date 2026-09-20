"""多格式文档解析：将上传文件转为纯文本。"""

from __future__ import annotations

import csv
import io
import json
import re
from html.parser import HTMLParser
from pathlib import Path


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self._parts: list[str] = []
        self._skip = False

    def handle_starttag(self, tag: str, attrs) -> None:  # noqa: ANN001
        if tag in {"script", "style"}:
            self._skip = True

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"}:
            self._skip = False

    def handle_data(self, data: str) -> None:
        if not self._skip:
            text = data.strip()
            if text:
                self._parts.append(text)

    def get_text(self) -> str:
        return "\n".join(self._parts)


def _decode_bytes(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gbk", "gb2312", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _load_plain_text(data: bytes) -> str:
    return _decode_bytes(data).strip()


def _load_html(data: bytes) -> str:
    parser = _HTMLTextExtractor()
    parser.feed(_decode_bytes(data))
    return parser.get_text().strip()


def _load_json(data: bytes) -> str:
    text = _decode_bytes(data)
    try:
        obj = json.loads(text)
        return json.dumps(obj, ensure_ascii=False, indent=2)
    except json.JSONDecodeError:
        return text.strip()


def _load_csv(data: bytes) -> str:
    text = _decode_bytes(data)
    reader = csv.reader(io.StringIO(text))
    rows = [", ".join(cell.strip() for cell in row if cell is not None) for row in reader]
    return "\n".join(row for row in rows if row).strip()


def _load_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            parts.append(page_text.strip())
    return "\n\n".join(parts).strip()


def _load_docx(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
    # 表格内容一并提取
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _load_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"## 工作表：{sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _load_pptx(data: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        slide_parts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text and shape.text.strip():
                slide_parts.append(shape.text.strip())
        if slide_parts:
            parts.append(f"## 幻灯片 {index}")
            parts.extend(slide_parts)
    return "\n".join(parts).strip()


def get_extension(filename: str) -> str:
    return Path(filename).suffix.lower().lstrip(".")


def extract_text_from_bytes(filename: str, data: bytes) -> str:
    """
    根据文件名后缀解析文件内容为纯文本。

    支持：txt/md/csv/json/html/pdf/docx/xlsx/pptx 等。
    """
    if not data:
        raise ValueError("文件内容为空")

    ext = get_extension(filename)
    loaders = {
        "txt": _load_plain_text,
        "md": _load_plain_text,
        "markdown": _load_plain_text,
        "log": _load_plain_text,
        "csv": _load_csv,
        "json": _load_json,
        "html": _load_html,
        "htm": _load_html,
        "pdf": _load_pdf,
        "docx": _load_docx,
        "xlsx": _load_xlsx,
        "pptx": _load_pptx,
    }

    # 老格式 .doc/.xls/.ppt 提示用户另存为新格式
    if ext in {"doc", "xls", "ppt"}:
        raise ValueError(
            f"暂不支持直接解析 .{ext}，请先另存为 "
            f".{'docx' if ext == 'doc' else 'xlsx' if ext == 'xls' else 'pptx'} 后再上传"
        )

    loader = loaders.get(ext)
    if loader is None:
        # 未知后缀：尝试按文本读取
        text = _load_plain_text(data)
        if not text.strip():
            raise ValueError(f"暂不支持的文件格式：.{ext or 'unknown'}")
        return text

    text = loader(data)
    text = re.sub(r"\n{3,}", "\n\n", (text or "").strip())
    if not text:
        raise ValueError("未能从文件中提取到有效文本")
    return text


def extract_text_from_upload(uploaded_file) -> str:
    """从 Streamlit UploadedFile 提取文本。"""
    return extract_text_from_bytes(uploaded_file.name, uploaded_file.getvalue())
